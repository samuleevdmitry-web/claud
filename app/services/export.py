"""Экспорт текущей выборки тендеров в Excel."""

from __future__ import annotations

import io
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from app.classifier import RELEVANCE_LABELS, Relevance
from app.config import get_settings
from app.services.tenders import TenderFilters, build_query
from app.sources.registry import BY_CODE
from app.sources.util import aware

COLUMNS = [
    ("Статус", 14),
    ("Score", 8),
    ("Срок подачи", 17),
    ("Название", 60),
    ("Заказчик", 45),
    ("ИНН", 13),
    ("Регион", 22),
    ("Место поставки", 35),
    ("Закон", 12),
    ("Способ", 22),
    ("НМЦК", 14),
    ("Валюта", 7),
    ("Опубликовано", 17),
    ("Номер ЕИС", 22),
    ("Категории", 30),
    ("Найденные ключи", 45),
    ("Маркеры", 30),
    ("Минус-слова", 30),
    ("Мой статус", 15),
    ("Заметки", 30),
    ("Площадки", 25),
    ("Ссылки", 60),
]
FILL = {
    "relevant": PatternFill("solid", fgColor="D6F0DF"),
    "review": PatternFill("solid", fgColor="FFF1D6"),
    "rejected": PatternFill("solid", fgColor="F9D9D9"),
}


def _labels(matches: list[dict]) -> dict[str, str]:
    out: dict[str, dict[str, None]] = {"product": {}, "marker": {}, "minus": {}}
    for m in matches:
        if m["kind"] in out and not m.get("suppressed"):
            out[m["kind"]][m["label"]] = None
    return {kind: ", ".join(values) for kind, values in out.items()}


def export_xlsx(session: Session, f: TenderFilters) -> bytes:
    tz = ZoneInfo(get_settings().timezone)

    def dt(value):
        value = aware(value)
        return value.astimezone(tz).replace(tzinfo=None) if value else None

    wb = Workbook()
    ws = wb.active
    ws.title = "Тендеры"
    ws.append([c[0] for c in COLUMNS])
    for i, (_, width) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
        ws.cell(1, i).font = Font(bold=True)
    ws.freeze_panes = "A2"

    for t in session.scalars(build_query(f)):
        labels = _labels(t.matches or [])
        ws.append(
            [
                RELEVANCE_LABELS.get(Relevance(t.relevance), t.relevance) if t.relevance else "",
                t.score,
                dt(t.application_deadline),
                t.title,
                t.customer_name,
                t.customer_inn,
                t.region,
                t.delivery_place,
                t.law,
                t.procedure_type,
                float(t.nmck) if t.nmck is not None else None,
                t.currency,
                dt(t.published_at),
                t.eis_number,
                ", ".join(t.categories or []),
                labels["product"],
                labels["marker"],
                labels["minus"],
                t.user_status,
                t.notes,
                ", ".join(
                    BY_CODE[s.source_code].title if s.source_code in BY_CODE else s.source_code
                    for s in t.sources
                ),
                "\n".join(s.url for s in t.sources if s.url),
            ]
        )
        row = ws.max_row
        if t.relevance in FILL:
            ws.cell(row, 1).fill = FILL[t.relevance]
        for col in (3, 13):
            ws.cell(row, col).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(row, 11).number_format = "# ##0.00"
        ws.cell(row, 22).alignment = Alignment(wrap_text=True)
    ws.auto_filter.ref = ws.dimensions
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
